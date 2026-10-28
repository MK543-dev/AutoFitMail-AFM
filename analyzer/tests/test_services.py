from io import BytesIO

from django.test import SimpleTestCase

from analyzer.services import jd_utils, scorer
from analyzer.services.parser import ParseError, extract_text


class ScorerTests(SimpleTestCase):
    def test_full_match(self):
        r = scorer.score('Python and Django developer', 'Required: Python, Django')
        self.assertEqual(r['score'], 100.0)

    def test_partial_match_and_missing(self):
        r = scorer.score('I know Python', 'We need Python, Django, Docker and AWS')
        self.assertEqual(r['score'], 25.0)
        self.assertEqual(r['missing'], ['AWS', 'Django', 'Docker'])

    def test_alias_matches(self):
        r = scorer.score('Worked with postgres daily', 'Needs PostgreSQL')
        self.assertEqual(r['score'], 100.0)

    def test_optional_skills_do_not_lower_score(self):
        r = scorer.score('Python', 'Python required. Nice to have: Docker')
        self.assertEqual(r['score'], 100.0)
        self.assertEqual(r['optional_missing'], ['Docker'])

    def test_no_required_skills_gives_zero(self):
        self.assertEqual(scorer.score('Python', 'Great team and culture')['score'], 0.0)


class JdUtilsTests(SimpleTestCase):
    def test_find_email(self):
        self.assertEqual(jd_utils.find_email('Send CV to hr@acme.com today'), 'hr@acme.com')

    def test_ignores_noreply(self):
        self.assertEqual(jd_utils.find_email('noreply@acme.com'), '')

    def test_required_optional_split(self):
        required, optional = jd_utils.split_required_optional('Python required. Nice to have: Docker')
        self.assertIn('Python', required)
        self.assertIn('Docker', optional)
        self.assertNotIn('Docker', required)

    def test_parse_jd_extracts_title_company_and_emails(self):
        jd = (
            '🚀 *Full Stack Engineer Intern | AI Platform | Kasparro*\n'
            '📍 *Location:* Bengaluru (On-site)\n'
            '🛠️ *Tech Stack:* Python, FastAPI, React, Docker, GitHub Actions\n'
            '📩 *Apply via Email:* nehakannan@kasparro.com\n'
            '📌 *CC:* grandmaster@kasparro.com\n'
        )
        parsed = jd_utils.parse_jd(jd)
        self.assertEqual(parsed['title'], 'Full Stack Engineer Intern')
        self.assertEqual(parsed['company'], 'Kasparro')
        self.assertEqual(parsed['apply_email'], 'nehakannan@kasparro.com')
        self.assertEqual(parsed['cc_emails'], ['grandmaster@kasparro.com'])

    def test_parse_jd_falls_back_to_email_domain_for_company(self):
        parsed = jd_utils.parse_jd('We need Python and Docker.\nSend your CV to hr@brightlabs.in')
        self.assertEqual(parsed['company'], 'Brightlabs')
        self.assertEqual(parsed['apply_email'], 'hr@brightlabs.in')

    def test_parse_jd_skips_generic_headline_for_company(self):
        jd = ('🚨 WE ARE HIRING | TECHNICAL SUPPORT ENGINEER 💻\n'
              'Role: Technical Support Engineer\n'
              'Send your resume to Muneendra@coretek.io')
        parsed = jd_utils.parse_jd(jd)
        self.assertEqual(parsed['title'], 'Technical Support Engineer')
        self.assertEqual(parsed['company'], 'Coretek')


class ParserTests(SimpleTestCase):
    def _docx_bytes(self):
        from docx import Document
        buf = BytesIO()
        doc = Document()
        doc.add_paragraph('Jane Doe - Python developer with Django and SQL experience.')
        doc.add_paragraph('Worked at Acme Corp from 2020 to 2024 building APIs and data pipelines.')
        doc.save(buf)
        buf.seek(0)
        return buf

    def test_docx_extraction(self):
        text = extract_text(self._docx_bytes(), 'cv.docx')
        self.assertIn('Django', text)

    def test_unsupported_extension(self):
        with self.assertRaises(ParseError):
            extract_text(BytesIO(b'x' * 200), 'cv.txt')

    def test_pdf_extraction(self):
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest('reportlab is not installed')
        buf = BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, 'Jane Doe - Python developer with Django, SQL and Git experience.')
        c.drawString(72, 700, 'Worked at Acme Corp from 2020 to 2024 building web applications and APIs.')
        c.drawString(72, 680, 'Looking for a backend role in a product company.')
        c.save()
        buf.seek(0)
        text = extract_text(buf, 'cv.pdf')
        self.assertIn('Django', text)

    def test_corrupt_pdf_gives_parse_error(self):
        with self.assertRaises(ParseError):
            extract_text(BytesIO(b'not a real pdf'), 'cv.pdf')
