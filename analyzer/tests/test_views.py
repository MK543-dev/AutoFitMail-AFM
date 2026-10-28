import io
import secrets

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from analyzer.models import Analysis, JobDescription, Resume


def _signup_password():
    """Generated at runtime so no placeholder credential lives in the repo."""
    return secrets.token_urlsafe(16)


def _docx_resume():
    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph('Python and Django developer with SQL and Git experience.')
    doc.add_paragraph('Building web applications and internal tools since 2020.')
    doc.save(buf)
    buf.seek(0)
    return buf.read()


class FlowSmokeTests(TestCase):
    def test_signup_upload_and_jd_extraction_flow(self):
        pw = _signup_password()
        resp = self.client.post(reverse('signup'), {
            'first_name': 'Test User',
            'email': 'test@example.com',
            'password1': pw,
            'password2': pw,
        })
        self.assertRedirects(resp, reverse('dashboard'))

        user = User.objects.get(username='test@example.com')
        upload = SimpleUploadedFile(
            'cv.docx', _docx_resume(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        resp = self.client.post(reverse('new_analysis'), {
            'new_resume': upload,
            'jd_text': ('Django Developer | Acme\n'
                        'We need Python, Django, Docker and AWS.\n'
                        'Apply: hr@acme.com'),
        })
        self.assertEqual(resp.status_code, 302)
        analysis = Analysis.objects.get()
        self.assertEqual(analysis.user, user)
        self.assertEqual(analysis.score, 50.0)
        self.assertEqual(sorted(analysis.matched_skills), ['Django', 'Python'])
        self.assertEqual(analysis.jd.apply_email, 'hr@acme.com')
        self.assertEqual(analysis.jd.company, 'Acme')
        self.assertIn('Django Developer', analysis.jd.title)

        resp = self.client.get(reverse('analysis_detail', args=[analysis.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Django Developer')

    def test_pages_require_login(self):
        for name in ('dashboard', 'new_analysis', 'tracker', 'resumes'):
            resp = self.client.get(reverse(name))
            self.assertEqual(resp.status_code, 302)

    def test_user_cannot_open_other_users_analysis(self):
        other = User.objects.create_user('b@x.com', 'b@x.com')
        viewer = User.objects.create_user('a@x.com', 'a@x.com')
        resume = Resume.objects.create(user=other, title='cv', text='Python ' * 30, file='resumes/2/x.txt')
        jd = JobDescription.objects.create(user=other, title='Dev', company='Acme', text='Python')
        analysis = Analysis.objects.create(user=other, resume=resume, jd=jd)
        self.client.force_login(viewer)
        resp = self.client.get(reverse('analysis_detail', args=[analysis.pk]))
        self.assertEqual(resp.status_code, 404)
