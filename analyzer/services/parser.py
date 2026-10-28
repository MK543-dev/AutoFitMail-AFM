import pdfplumber
from docx import Document


class ParseError(Exception):
    pass


def extract_text(file_obj, filename):
    name = filename.lower()
    if name.endswith('.pdf'):
        try:
            with pdfplumber.open(file_obj) as pdf:
                text = '\n'.join((page.extract_text() or '') for page in pdf.pages)
        except Exception as exc:
            raise ParseError(
                'Could not read this PDF. It may be scanned or damaged; '
                'upload a text PDF or DOCX.'
            ) from exc
    elif name.endswith('.docx'):
        doc = Document(file_obj)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
        text = '\n'.join(parts)
    else:
        raise ParseError('Only PDF and DOCX files are allowed.')
    text = text.strip()
    if len(text) < 100:
        raise ParseError(
            'No readable text found. Scanned PDFs are not supported; '
            'upload a text PDF or DOCX.'
        )
    return text
